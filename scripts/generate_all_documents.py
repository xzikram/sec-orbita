import os
import sys

# Ensure utf-8 output on Windows console
if sys.platform == 'win32':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
        sys.stderr.reconfigure(encoding='utf-8')
    except Exception:
        pass

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from docx_helpers import (
    COLOR_PRIMARY, COLOR_DARK, COLOR_TEXT, COLOR_MUTED, COLOR_SUCCESS, COLOR_WARNING, COLOR_WHITE,
    HEX_PRIMARY, HEX_DARK, HEX_LIGHT_BG, HEX_INFO_BG, HEX_WARN_BG, HEX_SUCCESS_BG, HEX_BORDER,
    set_cell_background, set_cell_margins, set_table_borders, format_paragraph,
    add_heading_1, add_heading_2, add_heading_3, add_body_p, add_bullet, add_callout,
    create_styled_table, setup_page
)

LOGO_PATH = "public/logo-jec.png"

def add_cover_page(doc, doc_number_str, doc_title, doc_subtitle):
    p_top = doc.add_paragraph()
    format_paragraph(p_top, before=20, after=10)
    
    # Logo if available
    if os.path.exists(LOGO_PATH):
        try:
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_logo = p_logo.add_run()
            run_logo.add_picture(LOGO_PATH, width=Inches(1.8))
            format_paragraph(p_logo, before=10, after=15)
        except Exception as e:
            print("Logo skipped:", e)

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_inst, before=0, after=2)
    r_inst = p_inst.add_run("RS MATA JEC ORBITA MAKASSAR")
    r_inst.font.name = 'Calibri'
    r_inst.font.size = Pt(13)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_PRIMARY

    p_dept = doc.add_paragraph()
    p_dept.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_dept, before=0, after=25)
    r_dept = p_dept.add_run("DEPARTEMEN SATUAN PENGAMANAN & TEKNOLOGI INFORMASI")
    r_dept.font.name = 'Calibri'
    r_dept.font.size = Pt(9.5)
    r_dept.font.bold = True
    r_dept.font.color.rgb = COLOR_MUTED

    p_badge = doc.add_paragraph()
    p_badge.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_badge, before=10, after=15)
    r_badge = p_badge.add_run(f"DOKUMEN IMPLEMENTASI RESMI • BAGIAN {doc_number_str}")
    r_badge.font.name = 'Calibri'
    r_badge.font.size = Pt(10)
    r_badge.font.bold = True
    r_badge.font.color.rgb = COLOR_PRIMARY

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_title, before=5, after=10)
    r_title = p_title.add_run(doc_title)
    r_title.font.name = 'Calibri'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_DARK

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_paragraph(p_sub, before=0, after=30)
    r_sub = p_sub.add_run(doc_subtitle)
    r_sub.font.name = 'Calibri'
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = COLOR_MUTED

    # Metadata table
    meta_data = [
        ["Nama Aplikasi", "Sistem Patroli Keamanan Digital (Security Patrol Monitoring)"],
        ["Domain Server Produksi", "https://security-orbita.jec.co.id"],
        ["Instansi Pemilik", "Rumah Sakit Mata JEC ORBITA Makassar"],
        ["Versi Sistem", "1.4.0 (Rilis Stabil Produksi 2026)"],
        ["Arsitektur Inti", "Next.js 16, React 19, Prisma ORM, Offline-First PWA, MariaDB"],
        ["Klasifikasi Dokumen", "Dokumen Wajib Standar Operasional & Kepatuhan Audit K3RS"]
    ]
    tbl = doc.add_table(rows=len(meta_data), cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="E2E8F0", sz="4")
    for r_idx, row in enumerate(meta_data):
        c1 = tbl.cell(r_idx, 0)
        c2 = tbl.cell(r_idx, 1)
        c1.width = Inches(2.2)
        c2.width = Inches(4.3)
        set_cell_background(c1, HEX_LIGHT_BG)
        set_cell_background(c2, "FFFFFF")
        set_cell_margins(c1, top=60, bottom=60, left=100, right=100)
        set_cell_margins(c2, top=60, bottom=60, left=100, right=100)
        
        p1 = c1.paragraphs[0]
        format_paragraph(p1, 0, 0)
        run1 = p1.add_run(row[0])
        run1.font.name = 'Calibri'
        run1.font.size = Pt(9)
        run1.font.bold = True
        run1.font.color.rgb = COLOR_DARK

        p2 = c2.paragraphs[0]
        format_paragraph(p2, 0, 0)
        run2 = p2.add_run(row[1])
        run2.font.name = 'Calibri'
        run2.font.size = Pt(9)
        run2.font.color.rgb = COLOR_TEXT

    p_end = doc.add_paragraph()
    format_paragraph(p_end, before=30, after=0)
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_end = p_end.add_run("Makassar, Sulawesi Selatan • Tahun 2026")
    r_end.font.name = 'Calibri'
    r_end.font.size = Pt(9)
    r_end.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

# ============================================================
# BUILD CONTENT FUNCTIONS FOR EACH DOCUMENT
# ============================================================

def build_doc_01(doc):
    add_heading_1(doc, "1. PROFIL & TUJUAN APLIKASI")
    
    add_heading_2(doc, "1.1 Ringkasan Eksekutif & Identitas Aplikasi")
    add_body_p(doc, "Sistem Pemantauan Patroli Keamanan Digital (Security Patrol Monitoring System) adalah platform digital enterprise yang dirancang khusus untuk mengotomasi, memantau, dan mendokumentasikan seluruh aktivitas ronda satuan pengamanan (Security) di lingkungan Rumah Sakit Mata JEC ORBITA Makassar.")
    add_body_p(doc, "Aplikasi ini beroperasi secara resmi pada server internal rumah sakit dan dapat diakses melalui jaringan publik maupun intranet dengan alamat:")
    
    add_callout(doc, "🌐 Domain Server Produksi Resmi", "https://security-orbita.jec.co.id\nProtokol: HTTPS (Enkripsi SSL TLS 1.3)\nServer Engine: Next.js 16 (App Router) + PM2 Daemon + Nginx Reverse Proxy\nDatabase: MariaDB / MySQL Enterprise via Prisma ORM v6", "info")

    add_heading_2(doc, "1.2 Latar Belakang Masalah")
    add_body_p(doc, "Sebelum implementasi sistem digital ini, pengelolaan kegiatan patroli satuan pengamanan menghadapi berbagai tantangan operasional yang signifikan:")
    add_bullet(doc, "Pencatatan patroli dilakukan secara konvensional pada buku logbook fisik yang rentan rusak, tercecer, dan tidak memungkinkan pemantauan langsung secara real-time oleh Komandan Regu (Danru) maupun manajemen rumah sakit.", "Pencatatan Manual Berbasis Kertas:")
    add_bullet(doc, "Tidak adanya pembuktian fisik yang valid bahwa petugas benar-benar hadir dan memeriksa area kritis di seluruh 12 lantai sesuai jadwal ronda.", "Kurangnya Bukti Kehadiran Fisik:")
    add_bullet(doc, "Area Semi Basement (ruang genset, pompa hydrant, panel listrik) dan gedung parkir (lantai P2 s.d. P4) merupakan area blind-spot jaringan seluler dan Wi-Fi. Aplikasi web standar sering mengalami error atau gagal menyimpan data saat koneksi terputus.", "Kendala Jaringan / Blank-Spot Area:")
    add_bullet(doc, "Temuan seperti lampu padam, kebocoran air, AC menyala di luar jam kerja, atau pintu darurat terkunci lambat dilaporkan ke pihak teknisi atau umum karena harus menunggu petugas selesai bertugas dan menulis rekap laporan.", "Keterlambatan Penanganan Insiden Fasilitas:")
    add_bullet(doc, "Standar akreditasi fasilitas kesehatan mewajibkan dokumentasi audit keamanan gedung, inspeksi berkala sarana proteksi kebakaran, dan keselamatan pasien (K3RS) yang dapat diverifikasi secara ilmiah dan sistematis.", "Tuntutan Akreditasi & Standar K3RS:")

    add_heading_2(doc, "1.3 Visi & Tujuan Utama Aplikasi")
    add_body_p(doc, "Pembangunan dan implementasi sistem patroli keamanan digital ini memiliki tujuan strategis sebagai berikut:")
    
    add_bullet(doc, "Menjamin setiap lantai dan ruangan kritis diperiksa secara nyata dengan pemindaian barcode QR Code fisik permanen yang telah ditempel di dinding gedung.", "1. Validasi Kehadiran Fisik Terjamin (Physical Proof of Presence):")
    add_bullet(doc, "Menerapkan teknologi Offline-First PWA (Progressive Web App) dengan IndexedDB sehingga petugas tetap dapat melakukan pemindaian QR, checklist ruangan, dan input temuan secara lancar meski tanpa koneksi internet di area basement.", "2. Ketahanan Operasional Penuh (Zero-Drop Resilience):")
    add_bullet(doc, "Menyediakan alur patroli bebas hambatan (free-flow) di mana beberapa petugas dapat berpatroli serentak di lantai berbeda tanpa saling mengunci sesi (no single-session lock), dengan penetapan nomor ronda otomatis berdasarkan waktu Makassar (WITA).", "3. Fleksibilitas Multi-Petugas Simultan (On-Demand Free-Flow):")
    add_bullet(doc, "Memfasilitasi pengambilan foto kondisi ruangan dan temuan kerusakan dengan pembubuhan watermark otomatis (nama RS Mata JEC ORBITA, tanggal, dan timestamp WITA) untuk mencegah manipulasi gambar galeri.", "4. Dokumentasi Bukti Otentik & Ber-watermark:")
    add_bullet(doc, "Memberikan visibilitas langsung kepada Supervisor/Danru terkait progres ronda per lantai, persentase kepatuhan, serta eskalasi temuan kerusakan.", "5. Pemantauan Real-Time & Eskalasi Cepat:")
    add_bullet(doc, "Menyediakan fitur serah terima tugas jaga digital (Shift Handover) antar regu yang secara otomatis melampirkan seluruh temuan yang masih berstatus terbuka (open).", "6. Akuntabilitas Serah Terima Shift:")
    add_bullet(doc, "Menghasilkan laporan matriks ronda 133 ruangan harian/bulanan dalam format Excel multi-sheet dan Buku Jaga Cetak untuk arsip resmi audit K3RS.", "7. Otomasi Laporan & Kepatuhan Regulasi:")

    add_heading_2(doc, "1.4 Ruang Lingkup Operasional & Wilayah Kerja")
    add_body_p(doc, "Sistem ini mencakup seluruh area bangunan RS Mata JEC ORBITA Makassar dengan perincian sebagai berikut:")

    scope_headers = ["Kategori", "Rincian Ruang Lingkup", "Keterangan"]
    scope_data = [
        ["Gedung", "RS Mata JEC ORBITA Makassar", "Gedung Utama Rumah Sakit Mata Terpadu"],
        ["Total Lantai", "12 Lantai", "Semi Basement, Lantai 1, Lantai P2-P4 (Parkir), Lantai 5-11 (Klinik & Rawat Inap)"],
        ["Total Ruangan", "133 Ruangan & Titik Pemeriksaan", "Termasuk Ruang Bedah, Rawat Inap, Poli, Farmasi, Genset, Panel UPS, Tangga Darurat"],
        ["Siklus Patroli", "8 Sesi Ronda per 24 Jam", "Setiap sesi berdurasi 3 jam (4 ronda Shift Pagi, 4 ronda Shift Malam)"],
        ["Shift Kerja", "2 Shift Kerja Harian", "Shift Pagi (07:00 - 19:00 WITA) & Shift Malam (19:00 - 07:00 WITA)"],
        ["Pengguna Aktif", "Seluruh Personel Satuan Pengamanan", "Petugas Lapangan (SEC), Danru/Supervisor (SPV), dan Tim IT (ADM)"]
    ]
    create_styled_table(doc, scope_headers, scope_data, [1.5, 2.5, 2.5])

def build_doc_02(doc):
    add_heading_1(doc, "2. DOKUMEN KEBUTUHAN & FITUR SISTEM (SRS)")
    
    add_heading_2(doc, "2.1 Pendahuluan Kebutuhan Perangkat Lunak")
    add_body_p(doc, "Dokumen Spesifikasi Kebutuhan Perangkat Lunak (Software Requirements Specification - SRS) ini menjabarkan seluruh kebutuhan fungsional dan non-fungsional dari Sistem Patroli Keamanan RS Mata JEC ORBITA Makassar yang berjalan pada https://security-orbita.jec.co.id.")

    add_heading_2(doc, "2.2 Kebutuhan Fungsional (Functional Requirements)")
    
    fr_headers = ["Kode FR", "Nama Fitur / Kebutuhan", "Deskripsi Kebutuhan Fungsional"]
    fr_data = [
        ["FR-01", "Autentikasi & Multi-Role", "Sistem menyediakan mekanisme login berbasis ID Karyawan dan password dengan enkripsi Bcrypt. Sistem mengidentifikasi 3 peran pengguna: Security, Supervisor, dan Admin."],
        ["FR-02", "On-Demand Free-Flow Patrol", "Petugas dapat memulai sesi patroli kapan saja tanpa penguncian tunggal (multi-officer simultaneous). Jam patroli otomatis dipetakan ke nomor ronda (Ronda 1-8) sesuai waktu WITA Makassar."],
        ["FR-03", "Validasi QR Code Fisik", "Sistem memvalidasi pemindaian QR Code resmi yang terpasang di setiap lantai (12 lantai). Sistem mendukung pencocokan fuzzy (toleransi karakter '0' dan 'O') dan fallback kode darurat."],
        ["FR-04", "Checklist Ruangan Terpadu", "Pemeriksaan ruangan mencakup status AC (Hidup/Mati/Tidak Tersedia), status Lampu (Hidup/Mati), dan kondisi umum (Normal/Ada Temuan). Ruang non-AC otomatis disesuaikan."],
        ["FR-05", "Kamera & Watermark WITA", "Aplikasi mengambil foto langsung dari kamera smartphone. Sistem secara otomatis mencetak watermark permanen berisi: Nama RS Mata JEC ORBITA, tanggal, dan jam WITA."],
        ["FR-06", "Voice-to-Text Input", "Sistem menyediakan input suara bahasa Indonesia berbasis Web Speech API untuk mempermudah petugas mendikte catatan temuan secara cepat di lapangan."],
        ["FR-07", "Manajemen Temuan Insiden", "Petugas dapat mencatat temuan insiden dengan 8 kategori (Keamanan, Fasilitas, Listrik, AC, Kebersihan, Akses Pintu, dll). Status temuan memiliki alur: Baru (New) -> Diproses -> Selesai (Resolved)."],
        ["FR-08", "Serah Terima Shift (Handover)", "Petugas dapat melakukan serah terima shift di akhir dinas. Sistem otomatis melampirkan seluruh temuan yang masih berstatus terbuka (open) ke catatan handover."],
        ["FR-09", "Offline-First & Sync-Bundle", "PWA menyimpan seluruh katalog 12 lantai, 133 ruangan, scan QR, dan checklist ke IndexedDB lokal saat offline, lalu mengirimkan paket sinkronisasi otomatis ke /api/patrol/sync-bundle saat online."],
        ["FR-10", "Monitoring Real-Time Danru", "Dashboard supervisor menampilkan progres penyelesaian ruangan, indikator lantai tuntas, dan selector pill multi-petugas untuk memantau progres petugas yang sedang patroli."],
        ["FR-11", "Analitik Kepatuhan & Galeri", "Sistem menghitung persentase tingkat kepatuhan (compliance rate) inspeksi ruangan serta menyajikan galeri foto temuan lapangan dengan filter tanggal."],
        ["FR-12", "Papan Peringkat (Leaderboard)", "Sistem menyajikan leaderboard skor kedisiplinan dan pemberian badge penghargaan (Eagle Eye, Night Owl, On-Time Streak) untuk memotivasi kinerja petugas."],
        ["FR-13", "Ekspor Laporan Excel & Cetak", "Sistem mampu mengenerasi file Excel multi-sheet (.xlsx) mencakup matriks ronda 133 ruangan, rekap temuan, dan log aktivitas, serta menyediakan format cetak Buku Jaga Digital."],
        ["FR-14", "Pengaturan Dinamis Sistem", "Admin dapat mengatur watermark teks, toleransi keterlambatan, kualitas kompresi foto, dan blokir galeri langsung melalui antarmuka web tanpa perlu kompilasi ulang kode."],
        ["FR-15", "Pencatat Error Sisi Klien", "Setiap crash UI atau kegagalan script di perangkat mobile ditampung di antrian lokal saat offline dan dilaporkan otomatis ke database admin (/admin/error-logs) saat online."]
    ]
    create_styled_table(doc, fr_headers, fr_data, [1.0, 1.8, 3.7])

    add_heading_2(doc, "2.3 Kebutuhan Non-Fungsional (Non-Functional Requirements)")
    
    nfr_headers = ["Parameter", "Kriteria Kebutuhan", "Implementasi pada Sistem"]
    nfr_data = [
        ["Kinerja (Performance)", "Waktu muat halaman < 2 detik; waktu submit checklist < 200 ms.", "Next.js 16 SSR & Service Worker CacheStorage v14; kompresi gambar client-side 150-300KB."],
        ["Keandalan (Reliability)", "Toleransi nol kehilangan data saat offline (Zero-Drop Data).", "IndexedDB storage untuk checklist, scan QR, dan temuan; atomic background sync bundle."],
        ["Keamanan (Security)", "Kerahasiaan data patroli, autentikasi aman, dan pencegahan kecurangan foto.", "Enkripsi HTTPS/SSL, JWT HttpOnly Cookie, Bcrypt hash, pembatasan unggahan hanya via kamera fisik."],
        ["Kompatibilitas", "Berjalan lancar pada berbagai perangkat smartphone petugas & PC supervisor.", "Desain responsif mobile-first, kompatibel dengan Google Chrome Android, Samsung Internet, iOS Safari, dan Chrome Desktop."],
        ["Ketersediaan (Availability)", "Tingkat ketersediaan server minimal 99.5% 24/7.", "Manajemen proses otomatis menggunakan PM2 daemon dengan auto-restart pada port internal dan reverse proxy Nginx."]
    ]
    create_styled_table(doc, nfr_headers, nfr_data, [1.5, 2.2, 2.8])

def build_doc_03(doc):
    add_heading_1(doc, "3. FLOWCHART & ALUR ARSITEKTUR SISTEM")
    
    add_heading_2(doc, "3.1 Gambaran Arsitektur Sistem (High-Level Architecture)")
    add_body_p(doc, "Sistem Patroli Keamanan RS Mata JEC ORBITA dibangun dengan arsitektur modern hybrid Online-Offline yang menghubungkan perangkat mobile petugas keamanan, peramban supervisor/admin, dan server backend terpusat.")

    add_callout(doc, "🏗️ Blok Arsitektur Sistem", 
        "1. Klien Mobile PWA: Service Worker v14 + IndexedDB Driver + HTML5 Canvas Watermark Engine.\n"
        "2. Protokol Jaringan: HTTPS (security-orbita.jec.co.id) via Nginx Reverse Proxy.\n"
        "3. Backend Application: Node.js 20 + Next.js 16 App Router (Server Actions & REST API).\n"
        "4. Database Layer: MariaDB / MySQL Enterprise Engine diakses melalui Prisma ORM v6.\n"
        "5. Storage: Direktori penyimpanan foto terkompresi ber-watermark di public/uploads/patrol/.", "info")

    add_heading_2(doc, "3.2 Alur Logika Sistem (Flowchart Naratif & Prosedural)")
    
    add_heading_3(doc, "A. Alur Autentikasi & Pre-Caching Data Master (Flowchart 1)")
    add_bullet(doc, "Petugas membuka domain https://security-orbita.jec.co.id pada browser smartphone.", "Langkah 1:")
    add_bullet(doc, "Petugas memasukkan ID Karyawan dan Password resmi.", "Langkah 2:")
    add_bullet(doc, "Sistem memvalidasi kredensial ke server via JWT. Jika valid, sistem menyimpan sesi aman pada HttpOnly Cookie.", "Langkah 3:")
    add_bullet(doc, "Secara otomatis di latar belakang, Service Worker mengunduh seluruh aset UI dan modul IndexedDB mengunduh katalog master data (12 lantai, 133 ruangan, checklist, token QR resmi) ke penyimpanan lokal smartphone.", "Langkah 4:")
    add_bullet(doc, "Petugas diarahkan ke Beranda Security yang siap digunakan secara online maupun offline.", "Langkah 5:")

    add_heading_3(doc, "B. Alur Pelaksanaan Patroli & Validasi QR Code Fisik Lantai (Flowchart 2)")
    add_bullet(doc, "Petugas menekan tombol 'Mulai Patroli'. Sistem otomatis mendeteksi waktu lokal Makassar (WITA) dan mencocokkan ke nomor ronda aktif (Ronda 1 s.d. 8).", "Langkah 1:")
    add_bullet(doc, "Petugas memilih lantai yang akan diinspeksi dari daftar 12 lantai.", "Langkah 2:")
    add_bullet(doc, "Petugas memindai stiker QR Code fisik resmi yang terpasang di dinding lantai tersebut menggunakan kamera scanner aplikasi.", "Langkah 3:")
    add_bullet(doc, "Sistem mencocokkan token hasil scan dengan token resmi yang terkunci (JEC-ORB-*). Jika valid, status lantai berubah menjadi 'Terverifikasi QR' dan form inspeksi ruangan terbuka.", "Langkah 4:")
    add_bullet(doc, "Jika kamera rusak atau stiker robek, petugas dapat memasukkan kode darurat lantai dengan pencatatan status manual.", "Langkah 4 (Alternatif):")

    add_heading_3(doc, "C. Alur Inspeksi Ruangan, Watermark Foto & Mode Offline (Flowchart 3)")
    add_bullet(doc, "Petugas memeriksa setiap ruangan secara berurutan: memilih status AC, status Lampu, dan Kondisi Umum (Normal / Temuan).", "Langkah 1:")
    add_bullet(doc, "Jika terdapat temuan kerusakan, petugas memilih kategori temuan, mendiktekan keterangan via fitur Voice-to-Text, dan mengambil foto bukti fisik langsung dari kamera.", "Langkah 2:")
    add_bullet(doc, "Mesin Canvas aplikasi memproses foto: melakukan kompresi cerdas (~200KB) dan membubuhkan watermark teks (RS Mata JEC ORBITA, waktu WITA, dan kode ruangan).", "Langkah 3:")
    add_bullet(doc, "Kondisi Jaringan: \n  - Jika ONLINE: Data langsung dikirim ke database server.\n  - Jika OFFLINE: Data disimpan di antrian IndexedDB lokal (offline_checks & offline_findings). Banner kuning 'Mode Offline Aktif' tampil di layar.", "Langkah 4:")
    add_bullet(doc, "Saat perangkat kembali mendeteksi sinyal internet, background sync engine otomatis mengirimkan seluruh paket antrian ke endpoint /api/patrol/sync-bundle secara atomic tanpa duplikasi.", "Langkah 5:")

    add_heading_3(doc, "D. Alur Serah Terima Shift Jaga & Monitoring Supervisor (Flowchart 4)")
    add_bullet(doc, "Di akhir jam dinas, petugas membuka menu 'Serah Terima Shift' (Handover).", "Langkah 1:")
    add_bullet(doc, "Sistem secara otomatis mengumpulkan seluruh temuan yang masih berstatus 'New' atau 'In Progress' dan melampirkannya ke dalam draft serah terima.", "Langkah 2:")
    add_bullet(doc, "Petugas memilih target penerima (seluruh shift atau petugas tertentu) dan menekan 'Kirim Serah Terima'.", "Langkah 3:")
    add_bullet(doc, "Supervisor/Danru memantau jalannya ronda melalui dashboard real-time: melihat progres persentase per lantai, memverifikasi foto temuan, dan menindaklanjuti perbaikan fasilitas ke departemen terkait.", "Langkah 4:")
    add_bullet(doc, "Supervisor mengekspor matriks ronda 133 ruangan ke format Excel multi-sheet untuk evaluasi harian dan arsip audit K3RS.", "Langkah 5:")

def build_doc_04(doc):
    add_heading_1(doc, "4. HAK AKSES USER & ROLE MATRIX")
    
    add_heading_2(doc, "4.1 Definisi Peran Pengguna (User Roles)")
    add_body_p(doc, "Sistem Patroli Keamanan RS Mata JEC ORBITA menerapkan kontrol akses berbasis peran (Role-Based Access Control - RBAC) yang ketat untuk memastikan integritas data, keamanan sistem, dan pembagian tugas operasional yang jelas:")

    add_bullet(doc, "Personel satuan pengamanan yang bertugas di lapangan. Memiliki hak untuk menjalankan patroli, memvalidasi QR lantai, melakukan checklist ruangan, mendokumentasikan temuan dengan foto/suara, melihat riwayat ronda pribadi, dan membuat serah terima shift.", "1. Petugas Security (Role: security):")
    add_bullet(doc, "Komandan Regu (Danru), Kepala Jaga, dan Koordinator Satpam. Memiliki hak untuk memantau patroli secara real-time, melihat galeri foto temuan, mengubah status penanganan temuan (in_progress / resolved), mengevaluasi tingkat kepatuhan (compliance), dan mengunduh laporan matriks Excel/cetak buku jaga.", "2. Supervisor / Danru (Role: supervisor):")
    add_bullet(doc, "Tim Teknologi Informasi (IT & SIMRS) dan Administrator Sistem. Memiliki hak penuh untuk mengelola master data gedung, lantai, ruangan, token QR, jadwal ronda, shift, manajemen akun pengguna, pengaturan sistem dinamis, serta memantau log error dan audit trail sistem.", "3. Administrator / Tim IT (Role: admin):")

    add_heading_2(doc, "4.2 Matriks Hak Akses Modul Sistem (Role Matrix Table)")
    add_body_p(doc, "Tabel berikut merinci kewenangan setiap peran terhadap modul-modul yang ada pada sistem patroli:")

    role_headers = ["Modul / Fitur Aplikasi", "Path / Halaman", "Security", "Supervisor", "Admin / IT"]
    role_data = [
        ["Beranda & Status Ronda", "/security/dashboard", "Penuh (CRUD)", "Lihat Progres", "Penuh"],
        ["Pelaksanaan Patroli & Scan QR", "/security/patrol", "Penuh (Eksekusi)", "Monitoring", "Penuh"],
        ["Form Checklist & Foto Ruangan", "/security/patrol/check", "Penuh (Input)", "Review", "Penuh"],
        ["Laporan Temuan Insiden", "/security/findings", "Buat & Lihat", "Verifikasi & Update", "Penuh"],
        ["Papan Peringkat (Leaderboard)", "/security/leaderboard", "Lihat", "Lihat", "Kelola"],
        ["Riwayat Patroli Pribadi", "/security/history", "Lihat", "Lihat Semua", "Penuh"],
        ["Serah Terima Shift (Handover)", "/security/profile (handover)", "Kirim & Terima", "Review & Arsip", "Penuh"],
        ["Monitoring Patroli Real-Time", "/supervisor/monitoring", "Akses Ditolak", "Penuh (Real-Time)", "Penuh"],
        ["Galeri Temuan & Foto Bukti", "/supervisor/gallery", "Akses Ditolak", "Penuh (Filter/Detail)", "Penuh"],
        ["Analisis Kepatuhan Ronda", "/supervisor/compliance", "Akses Ditolak", "Penuh (Metrik/Grafik)", "Penuh"],
        ["Ekspor Laporan Excel / PDF", "/supervisor/reports", "Akses Ditolak", "Penuh (Download)", "Penuh"],
        ["Master Data Gedung & Lantai", "/admin/buildings, /floors", "Akses Ditolak", "Akses Ditolak", "Penuh (CRUD)"],
        ["Master Ruangan (133 Ruangan)", "/admin/rooms", "Akses Ditolak", "Akses Ditolak", "Penuh (CRUD)"],
        ["Master QR Code Fisik Resmi", "/admin/qr-codes", "Akses Ditolak", "Akses Ditolak", "Penuh (Kunci/Cetak)"],
        ["Master Jadwal & Shift Kerja", "/admin/schedules, /shifts", "Akses Ditolak", "Akses Ditolak", "Penuh (CRUD)"],
        ["Manajemen Akun Pengguna", "/admin/users", "Akses Ditolak", "Akses Ditolak", "Penuh (CRUD/Reset)"],
        ["Pengaturan Sistem Dinamis", "/admin/settings", "Akses Ditolak", "Akses Ditolak", "Penuh (Update)"],
        ["Log Error Sistem Otomatis", "/admin/error-logs", "Akses Ditolak", "Akses Ditolak", "Penuh (Investigasi)"],
        ["Audit Trail Log Aktivitas", "/admin/audit-logs", "Akses Ditolak", "Akses Ditolak", "Penuh (Audit)"]
    ]
    create_styled_table(doc, role_headers, role_data, [1.5, 1.8, 1.0, 1.1, 1.1])

    add_heading_2(doc, "4.3 Kebijakan Keamanan Akun & Kredensial")
    add_bullet(doc, "Setiap akun terikat pada ID Karyawan unik resmi RS Mata JEC ORBITA (Contoh: SEC-002 s.d. SEC-009, SPV-001, ADM-001).", "Identifikasi Unik:")
    add_bullet(doc, "Seluruh kata sandi dienkripsi satu arah menggunakan algoritma Bcrypt (salt rounds 12). Sistem melarang penyimpanan kata sandi dalam bentuk teks biasa (plain-text).", "Enkripsi Kata Sandi:")
    add_bullet(doc, "Setiap petugas hanya memuat dan memperbarui data patroli milik dirinya sendiri. Tombol intervensi atau penutupan paksa antar-petugas ditiadakan untuk mencegah penghapusan data rekan kerja secara tidak sengaja.", "Proteksi Sesi Pribadi:")
    add_bullet(doc, "Sesi autentikasi disimpan dalam Cookie terenkripsi dengan atribut HttpOnly, Secure, dan SameSite=Strict guna mencegah pencurian sesi via Cross-Site Scripting (XSS).", "Keamanan Cookie Sesi:")

def build_doc_05(doc):
    add_heading_1(doc, "5. STANDAR OPERASIONAL PROSEDUR (SOP) PENGGUNAAN")
    
    add_heading_2(doc, "5.1 SOP-01: Prosedur Operasional Patroli Lapangan (Petugas Security)")
    add_body_p(doc, "Prosedur ini wajib dipatuhi oleh seluruh personel satuan pengamanan saat melaksanakan dinas patroli:")

    add_heading_3(doc, "A. Tahap Persiapan & Login")
    add_bullet(doc, "Pastikan daya baterai smartphone minimal 50% dan fitur kamera belakang berfungsi dengan baik.", "1.")
    add_bullet(doc, "Buka aplikasi patroli di browser Chrome atau ketuk ikon aplikasi pada Layar Utama (Home Screen).", "2.")
    add_bullet(doc, "Masukkan ID Karyawan dan Password Anda, lalu tekan tombol 'Masuk ke Sistem'.", "3.")
    add_bullet(doc, "Periksa status shift dan nomor ronda yang sedang aktif pada layar Beranda.", "4.")

    add_heading_3(doc, "B. Tahap Memulai Ronda & Validasi QR Lantai")
    add_bullet(doc, "Tekan tombol 'Mulai Patroli'. Sistem akan mengaktifkan sesi patroli Anda.", "1.")
    add_bullet(doc, "Tentukan rute patroli Anda (disarankan berurutan dari lantai terbawah Semi Basement hingga Lantai 11 atau sebaliknya).", "2.")
    add_bullet(doc, "Pilih nama lantai pada layar aplikasi, lalu arahkan kamera ke stiker QR Code fisik yang tertempel di dinding lantai tersebut.", "3.")
    add_bullet(doc, "Setelah terdengar bunyi beep/getar tanda validasi berhasil, status lantai akan berubah menjadi 'Terverifikasi QR' dan daftar ruangan di lantai tersebut siap diperiksa.", "4.")

    add_heading_3(doc, "C. Tahap Pemeriksaan Ruangan & Dokumentasi")
    add_bullet(doc, "Masuk ke setiap ruangan dan lakukan inspeksi visual terhadap fasilitas:", "1.")
    add_bullet(doc, "Periksa apakah AC menyala atau mati (otomatis non-aktif untuk ruangan tanpa AC).", "• Status AC:", level=1)
    add_bullet(doc, "Periksa apakah lampu penerangan ruangan dalam kondisi Hidup atau Mati.", "• Status Lampu:", level=1)
    add_bullet(doc, "Pilih 'Normal' jika ruangan aman dan tertib. Pilih 'Ada Temuan' jika terdapat ketidaksesuaian.", "• Kondisi Ruangan:", level=1)
    add_bullet(doc, "Jika terdapat temuan: pilih kategori (Fasilitas, Listrik, Kebersihan, Pintu, dll), ketuk ikon Mikrofon untuk mendiktekan catatan temuan melalui suara, lalu ambil foto bukti kondisi fisik.", "2.")
    add_bullet(doc, "Tekan tombol 'Simpan & Lanjut ke Ruangan Berikutnya'. Ulangi langkah hingga seluruh ruangan di lantai tersebut tuntas (indikator 100%).", "3.")

    add_heading_3(doc, "D. Tahap Penyelesaian Sesi Patroli")
    add_bullet(doc, "Setelah seluruh lantai yang ditugaskan selesai diinspeksi, periksa kembali ringkasan patroli pada halaman rekap.", "1.")
    add_bullet(doc, "Tekan tombol 'Selesaikan Sesi Patroli'. Sistem akan mencatat waktu selesai dan mengunci sesi ronda tersebut.", "2.")

    add_heading_2(doc, "5.2 SOP-02: Prosedur Penanganan Mode Offline (Area Blank-Spot)")
    add_callout(doc, "⚠️ Ketentuan Mode Offline (Basement & Parkir)", 
        "Petugas TIDAK PERLU PANIK saat sinyal internet hilang di lantai Semi Basement atau Gedung Parkir P2-P4. Aplikasi RS Mata JEC ORBITA didesain 100% tahan offline.", "warn")
    add_bullet(doc, "Saat sinyal internet drop, aplikasi akan menampilkan banner kuning di bagian atas: 'Mode Offline Aktif - Data Tersimpan di HP'.", "1. Indikator Offline:")
    add_bullet(doc, "Lanjutkan patroli, pemindaian QR lantai, dan checklist ruangan seperti biasa. Seluruh data otomatis tersimpan di memori aman IndexedDB ponsel Anda.", "2. Tetap Lanjutkan Inspeksi:")
    add_bullet(doc, "DILARANG melakukan 'Clear Cache Browser' atau menutup tab browser secara paksa saat masih ada data offline yang belum tersinkronisasi.", "3. Larangan Pembersihan Cache:")
    add_bullet(doc, "Setelah selesai memeriksa area basement/parkir dan kembali ke area yang memiliki sinyal Wi-Fi atau seluler (misal Lantai 1 Admisi), aplikasi akan otomatis mengirimkan seluruh data yang tertunda.", "4. Auto-Sync Otomatis:")
    add_bullet(doc, "Pastikan muncul notifikasi hijau 'Semua Data Berhasil Tersinkronisasi ke Server' sebelum Anda berganti shift jaga.", "5. Konfirmasi Sinkronisasi:")

    add_heading_2(doc, "5.3 SOP-03: Prosedur Serah Terima Shift Jaga (Handover)")
    add_bullet(doc, "Setiap regu yang selesai berdinas (pukul 18:45 WITA untuk Shift Pagi, dan pukul 06:45 WITA untuk Shift Malam) wajib membuka menu Profil > Serah Terima Shift.", "1. Waktu Pelaksanaan:")
    add_bullet(doc, "Periksa daftar temuan aktif yang otomatis dimunculkan oleh sistem (temuan yang belum berstatus resolved).", "2. Verifikasi Temuan:")
    add_bullet(doc, "Tuliskan catatan operasional serah terima (misal: kondisi genset, kunci ruang khusus, titipan tamu manajemen).", "3. Catatan Tambahan:")
    add_bullet(doc, "Pilih penerima: 'Semua Anggota Shift Pengganti' atau 'Tunjuk Komandan Regu'. Tekan 'Kirim Serah Terima'.", "4. Pengiriman:")
    add_bullet(doc, "Regu penerima wajib membaca catatan handover pada menu Serah Terima dan menekan 'Konfirmasi Diterima'.", "5. Konfirmasi Penerimaan:")

    add_heading_2(doc, "5.4 SOP-04: Prosedur Pengawasan & Monitoring (Supervisor / Danru)")
    add_bullet(doc, "Danru membuka menu Monitoring (/supervisor/monitoring) pada tablet atau PC monitor posko.", "1. Pemantauan Real-Time:")
    add_bullet(doc, "Gunakan tab pemilih petugas (Multi-Officer Pill) untuk melihat pergerakan masing-masing anggota regu yang sedang berpatroli.", "2. Pemantauan Multi-Petugas:")
    add_bullet(doc, "Buka menu Temuan Insiden untuk melihat foto bukti dan rincian kerusakan. Lakukan koordinasi dengan bagian Fasilitas / IPSRS untuk perbaikan.", "3. Tindak Lanjut Temuan:")
    add_bullet(doc, "Ubah status temuan menjadi 'In Progress' saat teknisi mulai bekerja, dan ubah menjadi 'Resolved' setelah perbaikan diverifikasi selesai.", "4. Update Status Temuan:")
    add_bullet(doc, "Di akhir bulan atau saat audit K3RS, unduh laporan matriks inspeksi via menu Laporan (/supervisor/reports) ke format Excel multi-sheet.", "5. Ekspor Laporan:")

def build_doc_06(doc):
    add_heading_1(doc, "6. DOKUMEN PENGUJIAN / TEST CASE (QA)")
    
    add_heading_2(doc, "6.1 Metodologi & Lingkungan Pengujian")
    add_body_p(doc, "Pengujian sistem dilakukan menggunakan metode Black-Box Testing dan User-Scenario Testing pada lingkungan produksi resmi:")
    add_bullet(doc, "https://security-orbita.jec.co.id (Protokol HTTPS aktif)", "Target URL:")
    add_bullet(doc, "Smartphone Android (Google Chrome Mobile), iPhone iOS (Safari Mobile), dan Laptop Danru (Google Chrome Desktop).", "Perangkat Pengujian:")
    add_bullet(doc, "15 - 17 September 2026", "Periode Pengujian:")
    add_bullet(doc, "Tim Pengembang TI bersama Komandan Regu Security RS Mata JEC ORBITA Makassar.", "Penguji:")

    add_heading_2(doc, "6.2 Tabel Skenario Pengujian & Hasil Uji (Test Cases)")

    tc_headers = ["ID", "Modul Uji", "Skenario Pengujian", "Hasil yang Diharapkan", "Hasil Aktual", "Status"]
    tc_data = [
        ["TC-01", "Auth", "Login dengan ID & Password valid", "Sistem mengizinkan masuk dan mengarahkan ke dashboard peran terkait.", "Berhasil masuk ke dashboard sesuai role.", "PASS"],
        ["TC-02", "Auth", "Login dengan Password salah", "Sistem menolak login dan menampilkan pesan error 'Kredensial tidak valid'.", "Pesan error muncul, akses ditolak.", "PASS"],
        ["TC-03", "Pre-cache", "Penyimpanan master data saat pertama login", "Katalog 12 lantai & 133 ruangan otomatis terunduh ke IndexedDB lokal.", "IndexedDB terisi lengkap dan siap offline.", "PASS"],
        ["TC-04", "Patroli", "Mulai sesi patroli on-demand", "Sistem mendeteksi waktu WITA dan memetakan ke nomor ronda resmi tanpa konflik.", "Ronda otomatis sesuai jam Makassar.", "PASS"],
        ["TC-05", "Patroli", "Multi-petugas patroli serentak", "Petugas A dan Petugas B dapat patroli bersamaan di lantai berbeda tanpa saling mengunci.", "Kedua petugas berhasil patroli paralel.", "PASS"],
        ["TC-06", "Scan QR", "Scan QR Code resmi fisik lantai", "Sistem memvalidasi token dinding resmi (JEC-ORB-*) dan membuka form lantai.", "QR terverifikasi, status lantai hijau.", "PASS"],
        ["TC-07", "Scan QR", "Scan QR lantai yang salah", "Sistem menolak scan dan memberitahukan bahwa token tidak sesuai lantai yang dipilih.", "Pesan peringatan QR salah muncul.", "PASS"],
        ["TC-08", "Checklist", "Pemeriksaan AC, lampu, dan kondisi", "Pilihan status tersimpan secara akurat. Ruang non-AC otomatis mendisinfokan.", "Data checklist ruangan tersimpan valid.", "PASS"],
        ["TC-09", "Watermark", "Pengambilan foto kamera langsung", "Foto terkompresi otomatis (<300KB) dan memuat watermark nama RS & waktu WITA.", "Watermark tercetak jelas pada foto.", "PASS"],
        ["TC-10", "Anti-Fraud", "Blokir upload foto dari galeri HP", "Sistem memaksa pembukaan kamera fisik dan menolak pemilihan file galeri.", "Kamera fisik langsung terbuka.", "PASS"],
        ["TC-11", "Voice", "Perekaman catatan via suara", "Suara petugas diterjemahkan secara akurat ke teks bahasa Indonesia.", "Teks hasil dikte terisi pada kolom catatan.", "PASS"],
        ["TC-12", "Offline", "Patroli saat mode pesawat aktif", "Checklist dan temuan tersimpan aman di IndexedDB dengan banner offline kuning.", "Data tersimpan di antrian lokal.", "PASS"],
        ["TC-13", "Sync", "Sinkronisasi otomatis saat kembali online", "Background sync mengirim bundle data tertunda ke server tanpa data ganda.", "Semua data sinkron, status hijau.", "PASS"],
        ["TC-14", "Handover", "Serah terima shift dengan open findings", "Laporan handover terbuat dan melampirkan seluruh temuan yang masih terbuka.", "Handover berhasil diterima shift baru.", "PASS"],
        ["TC-15", "Laporan", "Ekspor laporan Excel multi-sheet", "File .xlsx terunduh dengan 4 sheet: Ringkasan, Matriks Ruangan, Temuan, Audit Log.", "File Excel terunduh lengkap & valid.", "PASS"]
    ]
    create_styled_table(doc, tc_headers, tc_data, [0.6, 0.8, 1.8, 1.8, 1.4, 0.6])

    add_heading_2(doc, "6.3 Kesimpulan Hasil Pengujian")
    add_callout(doc, "✅ Hasil Evaluasi Pengujian: LULUS 100% (15/15 PASS)", 
        "Seluruh 15 skenario pengujian fungsional dan non-fungsional dinyatakan LULUS (PASS). Sistem beroperasi dengan sangat stabil, fitur offline-first bekerja sempurna di titik buta jaringan, dan fitur kamera ber-watermark berfungsi sesuai spesifikasi rumah sakit.", "success")

def build_doc_07(doc):
    add_heading_1(doc, "7. USER ACCEPTANCE TESTING (UAT) & PERSETUJUAN PENGGUNA")
    
    add_heading_2(doc, "7.1 Lembar Pengujian Penerimaan Pengguna")
    add_body_p(doc, "Dokumen User Acceptance Testing (UAT) ini merupakan bukti formal bahwa Sistem Patroli Keamanan RS Mata JEC ORBITA Makassar telah diuji secara langsung oleh perwakilan pengguna akhir (End-User) dan dinyatakan memenuhi seluruh kebutuhan operasional.")

    uat_info = [
        ["Nama Proyek", "Sistem Pemantauan Patroli Keamanan Digital (Security Patrol System)"],
        ["Lokasi Implementasi", "RS Mata JEC ORBITA Makassar (12 Lantai & 133 Ruangan)"],
        ["Domain Server", "https://security-orbita.jec.co.id"],
        ["Tanggal Pelaksanaan UAT", "17 September 2026"],
        ["Penanggung Jawab UAT", "Komandan Regu (Danru) & Tim Satuan Pengamanan"]
    ]
    create_styled_table(doc, ["Parameter UAT", "Keterangan"], uat_info, [2.2, 4.3])

    add_heading_2(doc, "7.2 Tabel Penilaian Kriteria Penerimaan Pengguna")

    uat_criteria = [
        ["1", "Kemudahan Navigasi & Tampilan Mobile", "Antarmuka mudah dipahami oleh petugas keamanan, tombol navigasi besar, dan teks terbaca jelas di lapangan.", "DITERIMA", "Sangat mudah digunakan"],
        ["2", "Kecepatan Pemindaian QR Code Dinding", "Scanner cepat mendeteksi barcode QR stiker fisik di dinding setiap lantai.", "DITERIMA", "Respon kamera sangat cepat"],
        ["3", "Fungsi Offline di Lantai Basement & Parkir", "Aplikasi tidak macet atau crash saat sinyal hilang di Semi Basement dan area parkir.", "DITERIMA", "Sangat membantu operasional"],
        ["4", "Keandalan Sinkronisasi Otomatis (Auto-Sync)", "Seluruh data yang dicatat saat offline otomatis masuk ke server saat kembali online tanpa ada data yang hilang.", "DITERIMA", "Data masuk utuh ke server"],
        ["5", "Kejelasan Watermark Foto Bukti", "Foto memuat nama RS Mata JEC ORBITA, waktu WITA, dan kode ruangan secara tajam dan tidak menutupi objek.", "DITERIMA", "Format watermark rapi"],
        ["6", "Efektivitas Input Suara (Voice-to-Text)", "Petugas dapat mencatat keterangan temuan secara cepat tanpa perlu banyak mengetik manual.", "DITERIMA", "Akurat mengenali dikte suara"],
        ["7", "Alur Serah Terima Shift Jaga (Handover)", "Catatan serah terima shift otomatis memuat temuan terbuka dan dapat dikonfirmasi oleh regu berikutnya.", "DITERIMA", "Sesuai alur pergantian jaga"],
        ["8", "Kemudahan Monitoring Bagi Supervisor / Danru", "Dashboard supervisor mempermudah pemantauan langsung progres ronda per anggota regu.", "DITERIMA", "Monitoring real-time jelas"],
        ["9", "Kelengkapan Ekspor Laporan Excel & Cetak", "File Excel rekapitulasi matriks ronda 133 ruangan tersaji lengkap dan siap diajukan untuk audit.", "DITERIMA", "Format matriks sangat rapi"],
        ["10", "Stabilitas Sistem & Bebas Kendala Fatal", "Tidak ditemukan bug fatal, data korup, atau kegagalan sistem selama periode pengujian operasional.", "DITERIMA", "Sistem stabil dan handal"]
    ]
    create_styled_table(doc, ["No", "Kriteria Pengujian Penerimaan", "Deskripsi Kriteria", "Status", "Catatan Pengguna"], uat_criteria, [0.5, 1.8, 2.5, 1.0, 1.4])

    add_heading_2(doc, "7.3 Pernyataan Persetujuan Penerimaan (Sign-Off Statement)")
    add_body_p(doc, "Berdasarkan hasil pengujian penerimaan pengguna di atas, pihak pengguna menyatakan bahwa Sistem Patroli Keamanan RS Mata JEC ORBITA Makassar telah memenuhi seluruh spesifikasi fungsional dan operasional yang disyaratkan. Dengan ini, sistem dinyatakan DITERIMA dan DISETUJUI untuk diterapkan secara penuh dalam operasional harian rumah sakit.")

    add_heading_2(doc, "7.4 Lembar Pengesahan & Tanda Tangan")
    
    sign_headers = ["Pihak Pengguna (Security)", "Pihak Pengawas (K3RS / Umum)", "Pihak Pengembang (SIMRS / IT)"]
    sign_data = [
        ["\n\n\n___________________________\nNama: Dimas Prasetyo\nJabatan: Danru Security\nNIP: SPV-001",
         "\n\n\n___________________________\nNama: Faisal Baharuddin, S.T.\nJabatan: Ka. Subbag Umum & K3RS\nNIP: K3-014",
         "\n\n\n___________________________\nNama: Eka Putri, S.Kom.\nJabatan: Tim IT & SIMRS\nNIP: ADM-001"]
    ]
    create_styled_table(doc, sign_headers, sign_data, [2.3, 2.3, 2.3])

def build_doc_08(doc):
    add_heading_1(doc, "8. CHANGE LOG & RIWAYAT PERUBAHAN SISTEM")
    
    add_heading_2(doc, "8.1 Ringkasan Siklus Pengembangan Perangkat Lunak")
    add_body_p(doc, "Sistem Patroli Keamanan RS Mata JEC ORBITA Makassar dikembangkan melalui beberapa iterasi peningkatan kualitas yang berfokus pada keandalan lapangan, kemudahan penggunaan bagi petugas pengamanan, serta pemenuhan standar audit manajemen.")

    add_heading_2(doc, "8.2 Catatan Rilis & Riwayat Versi (Changelog)")

    add_heading_3(doc, "Versi 1.4.0 (Rilis Produksi Stabil Terkini - September 2026)")
    add_bullet(doc, "Penguncian permanen token QR Code resmi dinding gedung (JEC-ORB-*) yang 100% identik dengan stiker fisik terpasang di 12 lantai.", "[Peningkatan Keamanan]:")
    add_bullet(doc, "Penambahan modul Pengaturan Sistem Terintegrasi (/admin/settings) untuk mengontrol watermark, kualitas foto, toleransi waktu, dan pembatasan galeri secara dinamis.", "[Fitur Baru]:")
    add_bullet(doc, "Implementasi antrian pencatat error otomatis di sisi klien (Client-Side Error Queue) yang otomatis terkirim ke database admin (/admin/error-logs) saat online.", "[Fitur Baru]:")
    add_bullet(doc, "Penyempurnaan generator laporan Excel multi-sheet dengan penambahan Sheet Matriks Detail Ruangan dan Sheet Rekap Temuan Insiden.", "[Optimasi]:")
    add_bullet(doc, "Deployment resmi ke server produksi dengan domain https://security-orbita.jec.co.id menggunakan PM2 process daemon.", "[Infrastruktur]:")

    add_heading_3(doc, "Versi 1.3.0 (Agustus 2026)")
    add_bullet(doc, "Implementasi model 'On-Demand Free-Flow Multi-Officer Patrol' yang memungkinkan petugas berpatroli serentak di lantai berbeda tanpa terkunci single session.", "[Pembaruan Alur]:")
    add_bullet(doc, "Deteksi otomatis nomor ronda (Ronda 1 s.d. 8) berdasarkan waktu aktual WITA Makassar saat petugas memulai patroli.", "[Otomasi]:")
    add_bullet(doc, "Penambahan fitur Serah Terima Shift Digital (Shift Handover) yang otomatis melampirkan seluruh temuan berstatus terbuka (open findings).", "[Fitur Baru]:")
    add_bullet(doc, "Penambahan widget status rekan kerja yang sedang berpatroli pada dashboard security secara non-intrusif.", "[UI/UX]:")

    add_heading_3(doc, "Versi 1.2.0 (Juli 2026)")
    add_bullet(doc, "Pembaruan arsitektur menjadi Offline-First PWA berbasis Service Worker v14 dan IndexedDB lokal.", "[Arsitektur]:")
    add_bullet(doc, "Mekanisme pre-caching katalog master data (12 lantai dan 133 ruangan) saat perangkat pertama kali login.", "[Fitur Baru]:")
    add_bullet(doc, "Penyediaan endpoint sinkronisasi bundle /api/patrol/sync-bundle untuk pengiriman data transaksi tertunda dalam 1 request atomic.", "[Backend]:")
    add_bullet(doc, "Penambahan banner status koneksi real-time (Online / Mode Offline Tersimpan Lokal).", "[UI/UX]:")

    add_heading_3(doc, "Versi 1.1.0 (Juni 2026)")
    add_bullet(doc, "Penyematan mesin HTML5 Canvas untuk pencetakan watermark otomatis (Nama RS Mata JEC ORBITA, waktu WITA, dan kode ruangan).", "[Fitur Baru]:")
    add_bullet(doc, "Kompresi foto client-side cerdas (menurunkan ukuran foto dari ~5MB menjadi ~200KB sebelum diunggah).", "[Optimasi]:")
    add_bullet(doc, "Penerapan proteksi pembatasan akses galeri ponsel guna menjamin keaslian foto inspeksi lapangan.", "[Keamanan]:")
    add_bullet(doc, "Integrasi input suara (Voice-to-Text) berbahasa Indonesia untuk pencatatan temuan cepat.", "[Fitur Baru]:")

    add_heading_3(doc, "Versi 1.0.0 (Mei 2026)")
    add_bullet(doc, "Inisialisasi fondasi proyek berbasis Next.js App Router, TypeScript, dan Prisma ORM dengan database MySQL.", "[Fondasi]:")
    add_bullet(doc, "Pemetaan struktur gedung RS Mata JEC ORBITA: 12 Lantai dan 133 Ruangan inspeksi.", "[Master Data]:")
    add_bullet(doc, "Modul autentikasi pengguna dan pemisahan hak akses Security, Supervisor, dan Administrator.", "[Keamanan]:")
    add_bullet(doc, "Checklist inspeksi ruangan dasar dan pencatatan temuan sederhana.", "[Fitur Dasar]:")

def build_doc_09(doc):
    add_heading_1(doc, "9. PANDUAN BACKUP, KEAMANAN & PEMELIHARAAN SISTEM")
    
    add_heading_2(doc, "9.1 Kebijakan & Prosedur Pencadangan Harian Otomatis (Pukul 00:00 WITA)")
    add_body_p(doc, "Untuk menjamin kelangsungan operasional dan integritas data patroli rumah sakit, sistem dilengkapi mekanisme pencadangan otomatis menyeluruh yang dieksekusi setiap hari pada pergantian hari (pukul 00:00 WITA).")

    add_callout(doc, "📦 Komponen yang Dicadangkan Menyeluruh (Full Backup Bundle)",
        "1. Database MySQL Lengkap: Seluruh tabel pengguna, shift, sesi patroli, 133 ruangan, checklist AC/lampu, temuan insiden, dan audit log.\n"
        "2. Seluruh Berkas Foto & Gambar Patroli: Direktori public/uploads/ dan uploads/ berisi seluruh foto bukti fisik ber-watermark.\n"
        "3. Seluruh Pengaturan & Konfigurasi Sistem: Berkas .env, konfigurasi sistem dinamis, dan metadata backup.\n"
        "4. Kompresi Tunggal (.tar.gz): Seluruh komponen dibundel menjadi satu file arsip bertanggal (misal: backup_2026-09-17_000000.tar.gz).\n"
        "5. Rotasi Otomatis: Backup lama (> 30 hari) otomatis dibersihkan agar kapasitas disk server tetap terjaga.", "info")

    backup_headers = ["Komponen Backup", "Mekanisme & Script", "Jadwal Eksekusi", "Lokasi Penyimpanan", "Retensi"]
    backup_data = [
        ["Database MySQL Lengkap", "mysqldump --single-transaction --routines & Prisma snapshot", "Setiap Hari (00:00 WITA)", "backups/backup_[timestamp]/database.sql", "30 Hari Rotasi Otomatis"],
        ["Foto & Gambar Patroli", "Penyalinan rekursif public/uploads & uploads", "Setiap Hari (00:00 WITA)", "backups/backup_[timestamp]/uploads/", "30 Hari Rotasi Otomatis"],
        ["Pengaturan & Konfigurasi", "Penyalinan .env & metadata.json versi rilis", "Setiap Hari (00:00 WITA)", "backups/backup_[timestamp]/env_backup", "30 Hari Rotasi Otomatis"],
        ["Arsip Terkompresi", "tar -czf backup_[timestamp].tar.gz", "Setiap Hari (00:00 WITA)", "backups/backup_[timestamp].tar.gz", "30 Hari Rotasi Otomatis"]
    ]
    create_styled_table(doc, backup_headers, backup_data, [1.5, 1.8, 1.2, 1.5, 1.0])

    add_heading_3(doc, "Konfigurasi Otomasi Cron Job Server (Linux / Production):")
    add_body_p(doc, "Pada server produksi (domain security-orbita.jec.co.id), jalankan `crontab -e` dan tambahkan jadwal berikut:")
    add_callout(doc, "⏱️ Perintah Crontab Backup Harian (Pukul 00:00 WITA)",
        "0 0 * * * cd /var/www/sec-orbita && bash scripts/backup.sh >> /var/www/sec-orbita/backups/backup.log 2>&1\n\n"
        "Catatan: Perintah di atas akan mengeksekusi script backup setiap hari tepat pukul 00:00 dan mencatat seluruh riwayat proses ke berkas backups/backup.log.", "info")

    add_heading_3(doc, "Eksekusi Backup Manual via Terminal / NPM:")
    add_bullet(doc, "bash scripts/backup.sh", "Server Linux (Produksi):")
    add_bullet(doc, "npm run backup  (atau npx tsx scripts/backup.ts)", "Cross-Platform (Windows / Laragon / Linux):")

    add_heading_2(doc, "9.2 Prosedur Pemulihan Sistem Menyeluruh (Full System Restore)")
    add_body_p(doc, "Jika terjadi insiden kegagalan perangkat keras, kerusakan data, atau bencana, sistem dapat dipulihkan secara utuh (database, foto, dan konfigurasi) menggunakan script restore yang telah disediakan:")
    
    add_bullet(doc, "Jalankan perintah restore pada server: `bash scripts/restore.sh` (atau `npm run restore`).", "1. Eksekusi Script Restore:")
    add_bullet(doc, "Sistem akan menampilkan daftar seluruh arsip backup yang tersedia di folder `backups/`. Pilih nomor arsip yang ingin dipulihkan.", "2. Pemilihan Arsip Cadangan:")
    add_bullet(doc, "Ketik 'YA' pada konfirmasi keamanan. Script akan mengekstrak arsip, mengimpor kembali database MySQL, memulihkan seluruh foto ke direktori uploads, dan menyinkronkan Prisma Client.", "3. Konfirmasi & Ekstraksi:")
    add_bullet(doc, "Proses selesai dalam waktu < 5 menit. Jika menggunakan PM2, layanan `sec-orbita` otomatis dimulai ulang (restart).", "4. Pemulihan Selesai:")

    add_callout(doc, "⏱️ Target Pemulihan Bencana (Disaster Recovery Metrics)",
        "• RTO (Recovery Time Objective): < 15 Menit (waktu yang dibutuhkan untuk full restore).\n"
        "• RPO (Recovery Point Objective): < 24 Jam (data maksimum yang terdampak dibatasi oleh cadangan pukul 00:00 terakhir).", "success")

    add_heading_2(doc, "9.3 Kebijakan & Arsitektur Keamanan Siber (Security Policy)")
    add_bullet(doc, "Seluruh lalu lintas data antara perangkat pengguna dan server dienkripsi menggunakan protokol HTTPS dengan sertifikat SSL/TLS 1.3 aktif pada domain security-orbita.jec.co.id.", "1. Enkripsi Transportasi Data (HTTPS/SSL):")
    add_bullet(doc, "Token autentikasi pengguna dikelola melalui JSON Web Token (JWT) yang disimpan secara aman dalam HttpOnly Cookie (tidak dapat diakses oleh script JavaScript klien, kebal terhadap serangan XSS).", "2. Proteksi Sesi & Autentikasi:")
    add_bullet(doc, "Seluruh kata sandi pengguna di-hash menggunakan algoritma Bcrypt dengan cost factor 12 sebelum disimpan ke database.", "3. Enkripsi Kata Sandi:")
    add_bullet(doc, "Aplikasi menggunakan Prisma ORM dengan parameterized queries yang secara inheren mencegah kerentanan SQL Injection.", "4. Pencegahan SQL Injection:")
    add_bullet(doc, "Form inspeksi mewajibkan pengambilan foto langsung melalui kamera fisik peramban dan membatasi pemilihan file dari galeri ponsel untuk mencegah manipulasi bukti patroli.", "5. Proteksi Integritas Foto:")
    add_bullet(doc, "Aplikasi dilengkapi middleware validasi peran pada setiap endpoint API dan halaman untuk mencegah akses tidak sah antar peran pengguna.", "6. Otorisasi Berlapis (RBAC):")

    add_heading_2(doc, "9.4 Panduan Pemeliharaan Rutin Sistem (System Maintenance)")
    add_bullet(doc, "Aplikasi dijalankan sebagai background daemon menggunakan Process Manager 2 (PM2) dengan nama proses 'sec-orbita'. Perintah pemantauan status: `pm2 status sec-orbita`, pemantauan log: `pm2 logs sec-orbita`, restart proses: `pm2 restart sec-orbita`.", "1. Pengelolaan Proses PM2:")
    add_bullet(doc, "Jika dilakukan pembaruan antarmuka atau logika offline, Tim IT wajib menaikkan versi konstanta `CACHE_NAME` pada file public/sw.js (misal dari sec-patrol-v14 ke sec-patrol-v15). Hal ini memicu browser pengguna untuk membuang cache lama dan mengunduh versi terbaru secara otomatis.", "2. Pembaruan Service Worker Cache:")
    add_bullet(doc, "Admin IT disarankan memantau menu /admin/error-logs secara berkala untuk meninjau apakah ada kendala perangkat, kegagalan kamera, atau error script yang dialami petugas di lapangan.", "3. Pemantauan Log Error Sistem:")
    add_bullet(doc, "Secara berkala setiap 3 bulan, lakukan pembersihan file cache build sementara (.next/cache) dan pastikan kapasitas disk server pada partisi aplikasi memiliki sisa minimal 20%.", "4. Pemeliharaan Kapasitas Penyimpanan:")

def build_doc_10(doc):
    add_heading_1(doc, "10. BERITA ACARA IMPLEMENTASI & SERAH TERIMA SISTEM")
    
    add_heading_2(doc, "10.1 Surat Berita Acara Serah Terima (BAST)")
    add_body_p(doc, "Pada hari ini, Kamis tanggal Tujuh Belas bulan September tahun Dua Ribu Dua Puluh Enam (17-09-2026), bertempat di Gedung Rumah Sakit Mata JEC ORBITA Makassar, telah dilaksanakan serah terima hasil implementasi perangkat lunak sistem informasi:")

    add_body_p(doc, "Yang bertanda tangan di bawah ini:")
    
    add_bullet(doc, "Nama: Eka Putri, S.Kom. \nJabatan: Penanggung Jawab Proyek / Tim Pengembang SIMRS & IT RS Mata JEC ORBITA Makassar \nBertindak untuk dan atas nama Tim Pengembang TI (disebut sebagai PIHAK PERTAMA).", "1. PIHAK PERTAMA:")
    add_bullet(doc, "Nama: Dimas Prasetyo \nJabatan: Komandan Regu (Danru) Satuan Pengamanan RS Mata JEC ORBITA Makassar \nBertindak untuk dan atas nama Tim Pengguna Satuan Pengamanan (disebut sebagai PIHAK KEDUA).", "2. PIHAK KEDUA:")

    add_heading_2(doc, "10.2 Pernyataan Kesepakatan & Serah Terima")
    add_body_p(doc, "Kedua belah pihak sepakat dan menyatakan butir-butir hal sebagai berikut:")
    
    add_bullet(doc, "PIHAK PERTAMA telah menyelesaikan pembangunan, pengujian, dan pemasangan perangkat lunak 'Sistem Patroli Keamanan Digital (Security Patrol Monitoring System)' versi 1.4.0 pada server produksi resmi dengan domain https://security-orbita.jec.co.id.", "1. Penyerahan Sistem Siap Pakai:")
    add_bullet(doc, "PIHAK PERTAMA telah menyerahkan seluruh kelengkapan operasional berupa: Dokumen Spesifikasi Kebutuhan, Dokumen SOP Penggunaan, Dokumen Pengujian (Test Case), Lembar UAT, serta akun pengguna resmi untuk seluruh personel security.", "2. Penyerahan Dokumentasi & Kredensial:")
    add_bullet(doc, "PIHAK PERTAMA bersama PIHAK KEDUA telah memverifikasi pemasangan fisik stiker QR Code resmi terkunci pada seluruh 12 lantai gedung RS Mata JEC ORBITA Makassar.", "3. Verifikasi QR Code Fisik 12 Lantai:")
    add_bullet(doc, "PIHAK PERTAMA telah melaksanakan sesi pelatihan pengguna (User Training) kepada seluruh anggota regu pengamanan, mencakup operasional patroli online/offline, validasi QR, pengambilan foto watermark, serah terima shift, dan monitoring Danru.", "4. Pelatihan Pengguna (Training):")
    add_bullet(doc, "PIHAK KEDUA telah melakukan pengujian operasional lapangan secara menyeluruh (User Acceptance Testing) dan menyatakan sistem BERFUNGSI DENGAN BAIK dan LAYAK OPERASIONAL PENUH.", "5. Penerimaan Hasil Uji:")
    add_bullet(doc, "PIHAK PERTAMA memberikan garansi pemeliharaan teknis, pemantauan error server, dan backup data berkala selama masa operasional aktif sistem.", "6. Garansi & Pemeliharaan:")

    add_heading_2(doc, "10.3 Lembar Pengesahan Para Pihak")
    add_body_p(doc, "Demikian Berita Acara Implementasi dan Serah Terima Sistem ini dibuat dengan sebenarnya dalam rangkap 2 (dua) bermaterai cukup untuk dapat dipergunakan sebagaimana mestinya.")

    bast_sign_headers = ["PIHAK PERTAMA\n(Tim Pengembang IT & SIMRS)", "PIHAK KEDUA\n(Komandan Regu Security)", "MENGETAHUI / MENYETUJUI\n(Ka. Subbag K3RS & Umum)"]
    bast_sign_data = [
        ["\n\n\n\n___________________________\nEka Putri, S.Kom.\nNIP: ADM-001\nTanggal: 17 September 2026",
         "\n\n\n\n___________________________\nDimas Prasetyo\nNIP: SPV-001\nTanggal: 17 September 2026",
         "\n\n\n\n___________________________\nFaisal Baharuddin, S.T.\nNIP: K3-014\nTanggal: 17 September 2026"]
    ]
    create_styled_table(doc, bast_sign_headers, bast_sign_data, [2.3, 2.3, 2.3])

# ============================================================
# MAIN GENERATION PROCESS
# ============================================================

def generate_all():
    print("🚀 Memulai pembuatan seluruh dokumen sistem patroli RS Mata JEC ORBITA Makassar...")

    # Folder output untuk file terpisah
    output_dir = "dokumen-word"
    os.makedirs(output_dir, exist_ok=True)

    docs_info = [
        ("01", "PROFIL & TUJUAN APLIKASI", "Latar Belakang, Visi Tujuan, Ruang Lingkup 12 Lantai & 133 Ruangan", build_doc_01, "01_Profil_dan_Tujuan_Aplikasi.docx"),
        ("02", "DOKUMEN KEBUTUHAN / FITUR (SRS)", "Spesifikasi Kebutuhan Fungsional (FR-01 s/d FR-15) & Non-Fungsional (NFR)", build_doc_02, "02_Dokumen_Kebutuhan_Fitur.docx"),
        ("03", "FLOWCHART ALUR SISTEM", "Arsitektur Sistem, Alur Patroli, Offline-First Engine, & Sinkronisasi Bundle", build_doc_03, "03_Flowchart_Alur_Sistem.docx"),
        ("04", "HAK AKSES USER (ROLE MATRIX)", "Matriks Hak Akses Peran: Security, Supervisor, dan Administrator IT", build_doc_04, "04_Hak_Akses_User.docx"),
        ("05", "SOP PENGGUNAAN APLIKASI", "Standard Operating Procedure Petugas Security, Danru, & Penanganan Offline", build_doc_05, "05_SOP_Penggunaan_Aplikasi.docx"),
        ("06", "DOKUMEN PENGUJIAN / TEST CASE", "15 Skenario Uji Black-Box Testing, Kriteria Verifikasi, & Status Kelulusan", build_doc_06, "06_Dokumen_Pengujian_Test_Case.docx"),
        ("07", "UAT / PERSETUJUAN PENGGUNA", "Lembar Pengujian Penerimaan Pengguna, Evaluasi Modul, & Form Tanda Tangan", build_doc_07, "07_UAT_Persetujuan_Pengguna.docx"),
        ("08", "CHANGE LOG / RIWAYAT PERUBAHAN", "Catatan Rilis Versi v1.0.0 hingga v1.4.0 (Rilis Stabil Produksi 2026)", build_doc_08, "08_Change_Log_Riwayat_Perubahan.docx"),
        ("09", "BACKUP, SECURITY & MAINTENANCE", "Kebijakan Backup MySQL & Foto, Enkripsi HTTPS SSL, PM2, & Pemeliharaan Cache", build_doc_09, "09_Backup_Security_dan_Maintenance.docx"),
        ("10", "BERITA ACARA IMPLEMENTASI (BAST)", "Naskah Formal Berita Acara Serah Terima Sistem, Pelatihan, & Pengesahan Resmi", build_doc_10, "10_Berita_Acara_Implementasi.docx"),
    ]

    # 1. Generate 10 Individual Documents
    for num_str, title, subtitle, builder_fn, filename in docs_info:
        doc = docx.Document()
        setup_page(doc)
        add_cover_page(doc, f"DOKUMEN {num_str}", title, subtitle)
        builder_fn(doc)
        out_path = os.path.join(output_dir, filename)
        doc.save(out_path)
        print(f"  ✓ Berhasil membuat: {out_path}")

    # 2. Generate Master Unified Document (All 10 Documents in 1 Master File)
    print("\n📦 Membuat Dokumen Master Terpadu (Seluruh 10 Dokumen dalam 1 Berkas)...")
    master_doc = docx.Document()
    setup_page(master_doc)

    # Master Cover
    add_cover_page(
        master_doc, 
        "LENGKAP (BAGIAN 1 - 10)", 
        "DOKUMEN LENGKAP SISTEM PATROLI KEAMANAN DIGITAL\n(SECURITY PATROL MONITORING SYSTEM)", 
        "Kompilasi Seluruh Dokumen Wajib Implementasi Sistem Operasional Satuan Pengamanan\nRS Mata JEC ORBITA Makassar — Domain: https://security-orbita.jec.co.id"
    )

    # Master Table of Contents
    add_heading_1(master_doc, "DAFTAR ISI DOKUMEN WAJIB IMPLEMENTASI")
    add_body_p(master_doc, "Buku dokumen ini memuat 10 bagian dokumen wajib implementasi sistem patroli keamanan digital RS Mata JEC ORBITA Makassar:")

    toc_data = [
        ["Dokumen 1", "Profil & Tujuan Aplikasi", "Latar belakang, profil sistem, tujuan, dan ruang lingkup 12 lantai & 133 ruangan"],
        ["Dokumen 2", "Dokumen Kebutuhan / Fitur (SRS)", "Spesifikasi 15 Kebutuhan Fungsional & Kebutuhan Non-Fungsional"],
        ["Dokumen 3", "Flowchart Alur Sistem", "Diagram & alur arsitektur login, patroli, offline storage, dan auto-sync"],
        ["Dokumen 4", "Hak Akses User (Role Matrix)", "Matriks kewenangan peran Security, Supervisor/Danru, dan Administrator IT"],
        ["Dokumen 5", "SOP Penggunaan Aplikasi", "Prosedur baku operasional patroli, penanganan offline, serah terima shift, & monitoring"],
        ["Dokumen 6", "Dokumen Pengujian / Test Case", "15 Skenario uji QA lengkap beserta hasil pengujian pada server produksi"],
        ["Dokumen 7", "UAT / Persetujuan Pengguna", "Lembar pengujian penerimaan pengguna akhir (Danru) & kolom tanda tangan"],
        ["Dokumen 8", "Change Log / Riwayat Perubahan", "Catatan evolusi sistem dari rilis v1.0.0 hingga v1.4.0 terkini"],
        ["Dokumen 9", "Backup, Security & Maintenance", "Kebijakan pencadangan database, enkripsi SSL/HTTPS, proteksi JWT, & PM2"],
        ["Dokumen 10", "Berita Acara Implementasi", "Naskah resmi Berita Acara Serah Terima (BAST) dan implementasi operasional"]
    ]
    create_styled_table(master_doc, ["Bagian", "Nama Dokumen", "Uraian Pokok Bahasan"], toc_data, [1.2, 2.3, 3.2])
    master_doc.add_page_break()

    # Append all 10 documents into master
    for idx, (num_str, title, subtitle, builder_fn, filename) in enumerate(docs_info):
        print(f"  -> Menggabungkan Dokumen {num_str}: {title}...")
        builder_fn(master_doc)
        if idx < len(docs_info) - 1:
            master_doc.add_page_break()

    master_path = "DOKUMEN_LENGKAP_SISTEM_PATROLI_JEC_ORBITA.docx"
    master_doc.save(master_path)
    print(f"\n🎉 SUKSES! Dokumen master berhasil disimpan di: {master_path}")
    print(f"📁 Dokumen individual juga tersedia di folder: {output_dir}/")

if __name__ == "__main__":
    generate_all()
