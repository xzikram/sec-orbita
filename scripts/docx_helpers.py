import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colors
COLOR_PRIMARY = RGBColor(2, 132, 199)      # #0284C7 Sky/JEC Blue
COLOR_DARK = RGBColor(15, 23, 42)          # #0F172A Slate Dark
COLOR_TEXT = RGBColor(30, 41, 59)          # #1E293B Text
COLOR_MUTED = RGBColor(100, 116, 139)      # #64748B Muted
COLOR_SUCCESS = RGBColor(22, 101, 52)      # #166534 Green
COLOR_WARNING = RGBColor(180, 83, 9)       # #B45309 Amber
COLOR_WHITE = RGBColor(255, 255, 255)

HEX_PRIMARY = "0284C7"
HEX_DARK = "0F172A"
HEX_LIGHT_BG = "F8FAFC"
HEX_INFO_BG = "F0F9FF"
HEX_WARN_BG = "FFFBEB"
HEX_SUCCESS_BG = "F0FDF4"
HEX_BORDER = "E2E8F0"

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="D1D5DB", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_paragraph(p, before=0, after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line_spacing

def add_heading_1(doc, text):
    h = doc.add_paragraph()
    format_paragraph(h, before=14, after=6)
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = COLOR_PRIMARY
    return h

def add_heading_2(doc, text):
    h = doc.add_paragraph()
    format_paragraph(h, before=10, after=4)
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = COLOR_DARK
    return h

def add_heading_3(doc, text):
    h = doc.add_paragraph()
    format_paragraph(h, before=6, after=2)
    r = h.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.color.rgb = COLOR_TEXT
    return h

def add_body_p(doc, text, bold=False, italic=False, color=COLOR_TEXT):
    p = doc.add_paragraph()
    format_paragraph(p, before=0, after=4)
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return p

def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, before=1, after=3)
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    if bold_prefix:
        r_pre = p.add_run(bold_prefix + " ")
        r_pre.font.name = 'Calibri'
        r_pre.font.size = Pt(10)
        r_pre.font.bold = True
        r_pre.font.color.rgb = COLOR_DARK
    r = p.add_run(text)
    r.font.name = 'Calibri'
    r.font.size = Pt(10)
    r.font.color.rgb = COLOR_TEXT
    return p

def add_callout(doc, title, text, style_type="info"):
    # Create 1x1 table
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = tbl.cell(0, 0)
    
    bg_hex = HEX_INFO_BG
    border_color = HEX_PRIMARY
    title_color = COLOR_PRIMARY
    if style_type == "warn":
        bg_hex = HEX_WARN_BG
        border_color = "F59E0B"
        title_color = COLOR_WARNING
    elif style_type == "success":
        bg_hex = HEX_SUCCESS_BG
        border_color = "10B981"
        title_color = COLOR_SUCCESS

    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
    
    # Left border only
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_color}"/>'
        f'<w:top w:val="none"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    format_paragraph(p, before=0, after=2)
    r_t = p.add_run(title + "\n")
    r_t.font.name = 'Calibri'
    r_t.font.size = Pt(10.5)
    r_t.font.bold = True
    r_t.font.color.rgb = title_color

    r_txt = p.add_run(text)
    r_txt.font.name = 'Calibri'
    r_txt.font.size = Pt(9.5)
    r_txt.font.color.rgb = COLOR_TEXT

    # spacer after table
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, before=0, after=4)

def create_styled_table(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color="CBD5E1", sz="4")

    # Header Row
    hdr_row = tbl.rows[0]
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        set_cell_background(cell, HEX_DARK)
        set_cell_margins(cell, top=120, bottom=120, left=120, right=120)
        p = cell.paragraphs[0]
        format_paragraph(p, before=0, after=0)
        r = p.add_run(heading)
        r.font.name = 'Calibri'
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = COLOR_WHITE

    # Data Rows
    for row_idx, row_data in enumerate(data):
        row = tbl.rows[row_idx + 1]
        bg_color = HEX_LIGHT_BG if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, val in enumerate(row_data):
            cell = row.cells[col_idx]
            set_cell_background(cell, bg_color)
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
            p = cell.paragraphs[0]
            format_paragraph(p, before=0, after=0)
            r = p.add_run(str(val))
            r.font.name = 'Calibri'
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_TEXT

    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    # spacer
    p_sp = doc.add_paragraph()
    format_paragraph(p_sp, before=0, after=6)
    return tbl

def setup_page(doc):
    for sec in doc.sections:
        sec.top_margin = Inches(0.8)
        sec.bottom_margin = Inches(0.8)
        sec.left_margin = Inches(0.8)
        sec.right_margin = Inches(0.8)
        sec.page_width = Inches(8.27)   # A4
        sec.page_height = Inches(11.69) # A4

        # Header
        header = sec.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("RS Mata JEC ORBITA Makassar  |  Sistem Patroli Keamanan (security-orbita.jec.co.id)")
        hrun.font.name = 'Calibri'
        hrun.font.size = Pt(8)
        hrun.font.color.rgb = COLOR_MUTED

        # Footer
        footer = sec.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Dokumentasi Resmi Sistem Operasional Patroli Keamanan Digital • RS Mata JEC ORBITA Makassar")
        frun.font.name = 'Calibri'
        frun.font.size = Pt(8)
        frun.font.color.rgb = COLOR_MUTED

print("Helper definitions loaded successfully.")
